"""Parser multi-formato: file → testo grezzo per chunking.

Ogni parser ritorna una lista di "documenti grezzi": stringhe di testo lunghe,
una per "unità logica" (es. una slide, un foglio Excel, una pagina PDF).
Il chunker poi le spezza in chunks di dimensione adatta agli embeddings.

Supporta: PDF, DOCX, PPTX, XLSX, MD, TXT, URL.
"""
from __future__ import annotations

import sys
from dataclasses import dataclass
from pathlib import Path


@dataclass
class RawDoc:
    text: str
    metadata: dict


# ─────────────────────────────────────────────────────────────
# PDF
# ─────────────────────────────────────────────────────────────
def parse_pdf(path: Path) -> list[RawDoc]:
    try:
        from pypdf import PdfReader
    except ImportError:
        raise RuntimeError("pypdf non installato. uv pip install -e .")

    reader = PdfReader(str(path))
    docs: list[RawDoc] = []
    for i, page in enumerate(reader.pages):
        text = (page.extract_text() or "").strip()
        if not text:
            continue
        docs.append(RawDoc(
            text=text,
            metadata={"source": path.name, "type": "pdf", "page": i + 1},
        ))
    return docs


# ─────────────────────────────────────────────────────────────
# DOCX
# ─────────────────────────────────────────────────────────────
def parse_docx(path: Path) -> list[RawDoc]:
    try:
        import docx  # python-docx
    except ImportError:
        raise RuntimeError("python-docx non installato. uv pip install -e .")

    d = docx.Document(str(path))
    # Concatena tutti i paragrafi non vuoti
    paragraphs = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    # Aggiunge anche il testo dalle tabelle
    for table in d.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)
    text = "\n\n".join(paragraphs)
    if not text:
        return []
    return [RawDoc(text=text, metadata={"source": path.name, "type": "docx"})]


# ─────────────────────────────────────────────────────────────
# PPTX
# ─────────────────────────────────────────────────────────────
def parse_pptx(path: Path) -> list[RawDoc]:
    try:
        from pptx import Presentation
    except ImportError:
        raise RuntimeError("python-pptx non installato. uv pip install -e .")

    prs = Presentation(str(path))
    docs: list[RawDoc] = []
    for i, slide in enumerate(prs.slides):
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs).strip()
                    if line:
                        parts.append(line)
        # Speaker notes
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"\n[NOTES] {notes}")

        text = "\n".join(parts).strip()
        if not text:
            continue
        docs.append(RawDoc(
            text=text,
            metadata={"source": path.name, "type": "pptx", "slide": i + 1},
        ))
    return docs


# ─────────────────────────────────────────────────────────────
# XLSX
# ─────────────────────────────────────────────────────────────
def parse_xlsx(path: Path) -> list[RawDoc]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise RuntimeError("openpyxl non installato. uv pip install -e .")

    wb = load_workbook(str(path), data_only=True, read_only=True)
    docs: list[RawDoc] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows: list[str] = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(c.strip() for c in cells):
                rows.append(" | ".join(cells))
        if not rows:
            continue
        text = f"# Sheet: {sheet_name}\n\n" + "\n".join(rows)
        docs.append(RawDoc(
            text=text,
            metadata={"source": path.name, "type": "xlsx", "sheet": sheet_name},
        ))
    wb.close()
    return docs


# ─────────────────────────────────────────────────────────────
# MD / TXT
# ─────────────────────────────────────────────────────────────
def parse_text(path: Path) -> list[RawDoc]:
    text = path.read_text(encoding="utf-8", errors="replace").strip()
    if not text:
        return []
    suffix = path.suffix.lower().lstrip(".")
    return [RawDoc(text=text, metadata={"source": path.name, "type": suffix or "txt"})]


# ─────────────────────────────────────────────────────────────
# URL (fetch + estrazione testo principale)
# ─────────────────────────────────────────────────────────────
def parse_url(url: str) -> list[RawDoc]:
    try:
        import trafilatura
    except ImportError:
        raise RuntimeError("trafilatura non installato. uv pip install -e .")

    downloaded = trafilatura.fetch_url(url)
    if not downloaded:
        print(f"[parsers] URL non scaricato: {url}", file=sys.stderr)
        return []
    text = trafilatura.extract(
        downloaded,
        include_comments=False,
        include_tables=True,
        favor_recall=True,
    )
    if not text:
        return []
    return [RawDoc(text=text.strip(), metadata={"source": url, "type": "url"})]


# ─────────────────────────────────────────────────────────────
# Dispatch
# ─────────────────────────────────────────────────────────────
def parse_file(path: Path) -> list[RawDoc]:
    """Auto-dispatch del parser in base all'estensione."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        return parse_pdf(path)
    if ext == ".docx":
        return parse_docx(path)
    if ext == ".pptx":
        return parse_pptx(path)
    if ext == ".xlsx":
        return parse_xlsx(path)
    if ext in (".md", ".markdown", ".txt"):
        return parse_text(path)
    # File _urls.txt: gestito a parte dall'indexer (non è un "vero" file di contenuto)
    raise ValueError(f"Estensione non supportata: {ext} ({path.name})")


SUPPORTED_EXTS = (".pdf", ".docx", ".pptx", ".xlsx", ".md", ".markdown", ".txt")
